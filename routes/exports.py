"""结果导出（PRD §10.4）：POST 创建 → GET 查询 → GET /download 代理下发。

同步生成（P0 数据量下毫秒级），创建即 READY；24h 后过期。
字节经对象存储（storage.py：local / firebase）落地；下载必须携带会话 Cookie
且通过对象级鉴权——不提供任何匿名或签名 URL，一律后端代理 send_file(BytesIO)。
XLSX 全部单元格按字符串写入，保留前导零；CSV 带 UTF-8 BOM。
"""
from __future__ import annotations

import csv
import io
from datetime import timedelta

from flask import Blueprint, request, send_file

from auth import audit, current_user
from errors import (
    invalid_argument,
    not_found,
    ok,
    service_unavailable,
    state_conflict,
)
from extensions import db
from models import Export, File, PageResult, ResultVersion
from routes.ocr import TERMINAL, _get_job_for_user
import storage

exports_bp = Blueprint("exports", __name__)

FORMATS = {"XLSX", "CSV", "TXT"}
EXPORT_TTL = timedelta(hours=24)


def _current_version(job_id: str) -> ResultVersion | None:
    return ResultVersion.query.filter_by(job_id=job_id, is_current=True).first()


def _load_pages(version_id: str) -> list[PageResult]:
    return (
        PageResult.query.filter_by(result_version_id=version_id)
        .order_by(PageResult.page_number.asc())
        .all()
    )


def _base_name(job) -> str:
    file = db.session.get(File, job.file_id)
    name = (file.original_name if file else "") or job.id
    stem = name.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    return stem or job.id


def _build_xlsx(pages: list[PageResult], out: io.BytesIO, missing_pages: list[int]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as e:  # 部署缺依赖时给出可诊断的 503
        raise service_unavailable("EXPORT_UNAVAILABLE", "error.export_unavailable") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "Tables"
    for page in pages:
        for t in page.tables or []:
            ws.append([f"Page {page.page_number} · {t.get('table_id', '')}"])
            col_count = int(t.get("column_count") or 0)
            if col_count:
                ws.append([f"col_{i + 1}" for i in range(col_count)])
            grid: dict[tuple[int, int], str] = {}
            for c in t.get("cells") or []:
                grid[(c.get("row_index", 0), c.get("column_index", 0))] = str(
                    c.get("current_text") or ""
                )
            for r in range(int(t.get("row_count") or 0)):
                ws.append([grid.get((r, c), "") for c in range(col_count)])
            ws.append([])
    if missing_pages:
        ws.append([f"Missing pages (OCR partial): {', '.join(map(str, missing_pages))}"])

    ws_text = wb.create_sheet("Text")
    ws_text.append(["page", "text"])
    for page in pages:
        ws_text.append([str(page.page_number), page.text or ""])

    ws_fields = wb.create_sheet("Fields")
    ws_fields.append(["page", "name", "value", "source"])
    for page in pages:
        for f in page.fields or []:
            ws_fields.append(
                [
                    str(page.page_number),
                    str(f.get("name", "")),
                    str(f.get("value", "")),
                    str(f.get("source", "")),
                ]
            )
    wb.save(out)


def _build_csv(pages: list[PageResult]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["page", "table_id", "row_index", "column_index", "text"])
    for page in pages:
        for t in page.tables or []:
            for c in t.get("cells") or []:
                writer.writerow(
                    [
                        page.page_number,
                        t.get("table_id", ""),
                        c.get("row_index", 0),
                        c.get("column_index", 0),
                        str(c.get("current_text") or ""),
                    ]
                )
    return buf.getvalue().encode("utf-8-sig")


def _build_txt(pages: list[PageResult]) -> bytes:
    parts: list[str] = []
    for page in pages:
        parts.append(f"===== Page {page.page_number} =====")
        parts.append(page.text or "")
    return "\n".join(parts).encode("utf-8")


@exports_bp.post("/ocr/jobs/<job_id>/exports")
def create_export(job_id: str):
    user = current_user()
    job = _get_job_for_user(job_id, user)
    if job.status not in TERMINAL:
        raise state_conflict("job.not_terminal")
    version = _current_version(job.id)
    if version is None:
        raise not_found()

    body = request.get_json(silent=True) or {}
    fmt = str(body.get("format") or "").upper()
    if fmt not in FORMATS:
        raise invalid_argument(field_errors=[{"field": "format", "reason": "invalid"}])

    from models import _now

    exp = Export(
        job_id=job.id,
        result_version_id=version.id,
        format=fmt,
        status="GENERATING",
        created_by_user_id=user.id,
        expires_at=_now() + EXPORT_TTL,
    )
    db.session.add(exp)
    db.session.flush()

    pages = _load_pages(version.id)
    try:
        if fmt == "XLSX":
            buf = io.BytesIO()
            _build_xlsx(pages, buf, list(job.missing_pages or []))
            data = buf.getvalue()
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif fmt == "CSV":
            data = _build_csv(pages)
            mime = "text/csv"
        else:
            data = _build_txt(pages)
            mime = "text/plain"
        exp.storage_path = storage.put_bytes("exports", f"{exp.id}.{fmt.lower()}", data, mime)
    except Exception as e:  # noqa: BLE001 生成失败落 FAILED，可重试
        exp.status = "FAILED"
        db.session.commit()
        raise service_unavailable("EXPORT_FAILED", "error.export_failed") from e

    exp.status = "READY"
    db.session.commit()
    audit("export.create", "export", exp.id, {"job_id": job.id, "format": fmt})
    return ok(exp.to_dict(), status=201)


def _get_export_for_user(export_id: str) -> Export:
    from models import _now

    exp = db.session.get(Export, export_id)
    if exp is None:
        raise not_found()
    _get_job_for_user(exp.job_id, current_user())  # 对象级鉴权，跨用户 404
    if exp.status == "READY" and exp.expires_at and exp.expires_at <= _now():
        exp.status = "EXPIRED"
        db.session.commit()
        raise state_conflict("export.expired")
    return exp


@exports_bp.get("/exports/<export_id>")
def get_export(export_id: str):
    return ok(_get_export_for_user(export_id).to_dict())


@exports_bp.get("/exports/<export_id>/download")
def download_export(export_id: str):
    exp = _get_export_for_user(export_id)
    if exp.status != "READY" or not exp.storage_path:
        raise state_conflict("export.not_ready")
    if not storage.exists(exp.storage_path):
        raise state_conflict("export.not_ready")
    from models import OcrJob

    job = db.session.get(OcrJob, exp.job_id)
    base = _base_name(job) if job else exp.job_id
    audit("export.download", "export", exp.id)
    return send_file(
        io.BytesIO(storage.get_bytes(exp.storage_path)),
        mimetype={
            "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "CSV": "text/csv",
            "TXT": "text/plain",
        }[exp.format],
        as_attachment=True,
        download_name=f"{base}-v{exp.result_version_id[:8]}.{exp.format.lower()}",
    )
